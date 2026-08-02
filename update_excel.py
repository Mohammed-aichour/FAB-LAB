import openpyxl

file_path = r'd:\Apps\Desktop\Downloads\Stage_3A\Anti\DOSSIER Tableau de criticité (A-B-C) - Plan de maintenance préventive complet -  Gammes de maintenance illustrées - FABLAB UNIVERSIAPOLIS.xlsx'

wb = openpyxl.load_workbook(file_path)
ws = wb['02_Arborescence_Reelle']

# Search for FL-065
updated = False
for row in range(1, ws.max_row + 1):
    val = ws.cell(row=row, column=1).value
    if val and str(val).strip() == 'FL-065':
        # Quantity column is E (column 5)
        # Check current value
        print("Found FL-065 at row:", row, "Current Qty:", ws.cell(row=row, column=5).value)
        # Set quantity to 0 if empty
        if ws.cell(row=row, column=5).value is None or str(ws.cell(row=row, column=5).value).strip() == '':
            ws.cell(row=row, column=5, value=0)
            ws.cell(row=row, column=6, value='En rupture / A approvisionner')
            updated = True
            print("Updated FL-065 quantity to 0 and state to En rupture / A approvisionner")

if updated:
    wb.save(file_path)
    print("Excel file saved successfully!")
else:
    print("No changes needed in Excel file.")
