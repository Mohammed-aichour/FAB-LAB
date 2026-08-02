import openpyxl
import os

excel_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\GMAO_FabLab_Universiapolis_Dossier_complet_8-9-12-13-14.xlsx"

if os.path.exists(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    print("Sheets in Excel file:", wb.sheetnames)
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        print(f"Sheet '{sheetname}': {sheet.max_row} rows, {sheet.max_column} cols")
else:
    print("Excel file not found!")
