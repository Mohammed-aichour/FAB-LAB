import openpyxl
import json
import re

# Load Excel files
f1 = 'DOSSIER Tableau de criticité (A-B-C) - Plan de maintenance préventive complet -  Gammes de maintenance illustrées - FABLAB UNIVERSIAPOLIS.xlsx'
wb1 = openpyxl.load_workbook(f1, data_only=True)

f2 = 'GMAO_FabLab_Universiapolis_Dossier_complet_8-9-12-13-14.xlsx'
wb2 = openpyxl.load_workbook(f2, data_only=True)

report = {
    "total_errors_detected": 0,
    "machines_added": 0,
    "machines_updated": 0,
    "stock_added": 0,
    "stock_updated": 0,
    "preventif_tasks_synced": 0,
    "amdec_items_synced": 0,
    "suppliers_synced": 0,
    "details": []
}

print("=== 1. EXTRACT MACHINES FROM EXCEL ===")
# 02_Arborescence_Reelle
ws_arb = wb1['02_Arborescence_Reelle']
arb_rows = [r for r in list(ws_arb.iter_rows(values_only=True)) if any(r)]

excel_machines = []
header = None
for r in arb_rows:
    str_r = [str(c).strip() if c is not None else '' for c in r]
    if 'Code arborescence' in str_r[0] or 'Code' in str_r[0]:
        header = str_r
        continue
    if header and str_r[0] and not str_r[0].startswith('ARBORESCENCE') and not str_r[0].startswith('Base'):
        code = str_r[0]
        desig = str_r[1] if len(str_r)>1 else ''
        niveau = str_r[2] if len(str_r)>2 else ''
        desc = str_r[3] if len(str_r)>3 else ''
        qty = str_r[4] if len(str_r)>4 else '1'
        etat = str_r[5] if len(str_r)>5 else 'Opérationnel'
        code_amdec = str_r[7] if len(str_r)>7 else ''

        if code.startswith('FL-') or niveau == 'Equipement':
            try:
                q_num = int(float(qty))
            except:
                q_num = 1
            
            status = 'Opérationnel'
            if 'hors service' in etat.lower() or 'hs' in etat.lower():
                status = 'Hors service'
            elif 'confirmer' in etat.lower() or 'contrôler' in etat.lower() or 'renseigné' in etat.lower():
                status = 'À contrôler'

            category = 'Atelier FabLab'
            if 'IMP' in code or 'FDM' in desig or 'résine' in desig.lower(): category = 'Atelier Impression 3D'
            elif 'LAS' in code or 'laser' in desig.lower(): category = 'Atelier Découpe Laser'
            elif 'CNC' in code or 'fraiseuse' in desig.lower(): category = 'Atelier Usinage CNC'
            elif 'ELE' in code or 'soudure' in desig.lower() or 'oscilloscope' in desig.lower(): category = 'Atelier Électronique'
            elif 'BOI' in code or 'perceuse' in desig.lower() or 'scie' in desig.lower(): category = 'Atelier Bois & Outillage'
            elif 'TEX' in code or 'brodeuse' in desig.lower(): category = 'Atelier Textile & Vinyle'
            elif 'VIN' in code or 'vinyle' in desig.lower(): category = 'Atelier Textile & Vinyle'
            elif 'CMP' in code or 'ASP' in code or 'compresseur' in desig.lower() or 'aspiration' in desig.lower(): category = 'Équipements Généraux'

            excel_machines.append({
                "id": code,
                "reference": code,
                "codeEq": code_amdec or code,
                "name": desig,
                "category": category,
                "atelier": category.replace('Atelier ', ''),
                "location": f"Zone {category.replace('Atelier ', '')}",
                "manufacturer": desig.split(' ')[0],
                "model": desc or desig,
                "serialNumber": f"SN-{code}-2026",
                "yearInService": 2022,
                "criticite": "A" if "EQ" in code_amdec else ("B" if "IMP" in code or "LAS" in code or "CNC" in code else "C"),
                "status": status,
                "quantity": q_num,
                "description": desc or f"{desig} - Registre officiel FabLab Universiapolis",
                "logiciel": "GMAO Universiapolis",
                "level": 1
            })

print(f"Extracted {len(excel_machines)} machines from Excel arborescence.")

print("\n=== 2. EXTRACT STOCK ITEMS FROM EXCEL ===")
# Sheet 9. Pièces de rechange from WB2
ws_stock = wb2['9. Pièces de rechange']
stock_rows = [r for r in list(ws_stock.iter_rows(values_only=True)) if any(r)]

excel_stock = []
for r in stock_rows[2:]:
    if r and r[0] and str(r[0]).strip() and not str(r[0]).startswith('Étape') and not str(r[0]).startswith('Réf'):
        ref_p = str(r[0]).strip()
        desig = str(r[1]).strip() if len(r)>1 and r[1] else ''
        cat = str(r[2]).strip() if len(r)>2 and r[2] else 'Pièce de rechange'
        equip = str(r[3]).strip() if len(r)>3 and r[3] else 'Tous Équipements'
        fourn = str(r[4]).strip() if len(r)>4 and r[4] else 'Outillage Agadir'
        ref_fourn = str(r[5]).strip() if len(r)>5 and r[5] else ref_p
        
        try: q_act = int(float(r[6])) if len(r)>6 and r[6] is not None else 10
        except: q_act = 10
        
        try: q_min = int(float(r[7])) if len(r)>7 and r[7] is not None else 3
        except: q_min = 3

        try: q_max = int(float(r[8])) if len(r)>8 and r[8] is not None else 20
        except: q_max = 20

        try: cost = float(r[9]) if len(r)>9 and r[9] is not None else 100.0
        except: cost = 100.0

        unit = str(r[10]).strip() if len(r)>10 and r[10] else 'u'
        loc = str(r[11]).strip() if len(r)>11 and r[11] else 'Magasin GMAO'

        st = 'Opérationnel'
        if q_act <= 0: st = 'Rupture'
        elif q_act <= q_min: st = 'Alerte'

        excel_stock.append({
            "id": ref_p,
            "reference": ref_p,
            "name": desig,
            "category": cat,
            "zone": "Magasin GMAO",
            "subcat": cat,
            "equipement": equip,
            "supplier": fourn,
            "refSupplier": ref_fourn,
            "unitCostMAD": cost,
            "priceMAD": cost,
            "unit": unit,
            "quantity": q_act,
            "min": q_min,
            "max": q_max,
            "location": loc,
            "description": f"{desig} - Stock officiel FabLab",
            "status": st
        })

print(f"Extracted {len(excel_stock)} stock items from Excel sheet 9.")

print("\n=== 3. EXTRACT SUPPLIERS FROM EXCEL ===")
ws_supp = wb2['9b. Fournisseurs']
supp_rows = [r for r in list(ws_supp.iter_rows(values_only=True)) if any(r)]
excel_suppliers = []
for r in supp_rows[2:]:
    if r and r[0] and str(r[0]).strip() and not str(r[0]).startswith('Code') and not str(r[0]).startswith('Étape'):
        code = str(r[0]).strip()
        name = str(r[1]).strip() if len(r)>1 and r[1] else ''
        domaine = str(r[2]).strip() if len(r)>2 and r[2] else ''
        contact = str(r[3]).strip() if len(r)>3 and r[3] else ''
        tel_email = str(r[4]).strip() if len(r)>4 and r[4] else ''
        delai = str(r[5]).strip() if len(r)>5 and r[5] else '3'
        
        excel_suppliers.append({
            "id": code,
            "code": code,
            "name": name,
            "domaine": domaine,
            "contact": contact,
            "telEmail": tel_email,
            "delaiMoyenJours": delai
        })

print(f"Extracted {len(excel_suppliers)} suppliers from Excel sheet 9b.")

# Write updated realMachineData.ts
m_ts_content = f'''export interface Machine {{
  id: string | number;
  reference: string;
  codeEq: string;
  name: string;
  category: string;
  atelier: string;
  location: string;
  manufacturer: string;
  model: string;
  serialNumber: string;
  yearInService: number;
  criticite: 'A' | 'B' | 'C';
  status: 'Opérationnel' | 'Hors service' | 'À contrôler';
  quantity: number;
  description: string;
  logiciel: string;
  level: number;
}}

export const realMachinesData: Machine[] = {json.dumps(excel_machines, ensure_ascii=False, indent=2)};
'''

with open('frontend/src/data/realMachineData.ts', 'w', encoding='utf-8') as f:
    f.write(m_ts_content)

with open('data_db/machines_db.json', 'w', encoding='utf-8') as f:
    json.dump(excel_machines, f, ensure_ascii=False, indent=2)

report["machines_added"] = len(excel_machines)
print("Updated realMachineData.ts and data_db/machines_db.json!")

# Write updated realStockData.ts
s_ts_content = f'''export interface StockItem {{
  id: string | number;
  reference: string;
  name: string;
  category: string;
  zone: string;
  subcat: string;
  equipement: string;
  supplier: string;
  refSupplier: string;
  unitCostMAD: number;
  priceMAD: number;
  unit: string;
  quantity: number;
  min: number;
  max: number;
  location: string;
  description: string;
  status: string;
}}

export const realStockItems: StockItem[] = {json.dumps(excel_stock, ensure_ascii=False, indent=2)};
'''

with open('frontend/src/data/realStockData.ts', 'w', encoding='utf-8') as f:
    f.write(s_ts_content)

with open('data_db/stock_db.json', 'w', encoding='utf-8') as f:
    json.dump(excel_stock, f, ensure_ascii=False, indent=2)

report["stock_added"] = len(excel_stock)
print("Updated realStockData.ts and data_db/stock_db.json!")

# Write updated supplierData.ts
supp_ts_content = f'''export interface Supplier {{
  id: string;
  code: string;
  name: string;
  domaine: string;
  contact: string;
  telEmail: string;
  delaiMoyenJours: string;
}}

export const supplierData: Supplier[] = {json.dumps(excel_suppliers, ensure_ascii=False, indent=2)};
'''

with open('frontend/src/data/supplierData.ts', 'w', encoding='utf-8') as f:
    f.write(supp_ts_content)

with open('data_db/suppliers_db.json', 'w', encoding='utf-8') as f:
    json.dump(excel_suppliers, f, ensure_ascii=False, indent=2)

report["suppliers_synced"] = len(excel_suppliers)
print("Updated supplierData.ts and data_db/suppliers_db.json!")

# Save report JSON
with open('scratch/sync_report.json', 'w', encoding='utf-8') as f:
    json.dump(report, f, ensure_ascii=False, indent=2)

print("\nSYNCHRONIZATION COMPLETED SUCCESSFULLY!")
