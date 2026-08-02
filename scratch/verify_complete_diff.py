import openpyxl
import json
import re

excel_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\GMAO_FabLab_Universiapolis_Dossier_complet_8-9-12-13-14.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Equipments Sheet (Réf. Équipements)
sheet_eq = wb.worksheets[1]
excel_machines = {}
for r in range(5, 19):
    eq_id = sheet_eq.cell(r, 1).value
    if eq_id:
        excel_machines[eq_id] = {
            'id': eq_id,
            'name': sheet_eq.cell(r, 2).value,
            'atelier': sheet_eq.cell(r, 3).value,
            'manufacturer': sheet_eq.cell(r, 4).value,
            'criticite': sheet_eq.cell(r, 5).value,
        }

# Read realMachineData.ts
with open(r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\frontend\src\data\realMachineData.ts", "r", encoding="utf-8") as f:
    m_code = f.read()

print("=== 1. MACHINES VERIFICATION ===")
print("Excel Machines count:", len(excel_machines))

# 2. Stock Sheet (9. Pièces de rechange)
sheet_stk = wb.worksheets[4]
excel_stock = {}
for r in range(5, 35):
    ref = sheet_stk.cell(r, 1).value
    if ref:
        excel_stock[ref] = {
            'ref': ref,
            'name': sheet_stk.cell(r, 2).value,
            'category': sheet_stk.cell(r, 3).value,
            'equipement': sheet_stk.cell(r, 4).value,
            'supplier': sheet_stk.cell(r, 5).value,
            'price': float(sheet_stk.cell(r, 7).value or 0),
            'unit': sheet_stk.cell(r, 8).value,
            'stock': float(sheet_stk.cell(r, 9).value or 0),
            'minStock': float(sheet_stk.cell(r, 10).value or 0),
            'maxStock': float(sheet_stk.cell(r, 11).value or 0),
            'toOrder': float(sheet_stk.cell(r, 12).value or 0),
            'location': sheet_stk.cell(r, 14).value
        }

print("\n=== 2. STOCK VERIFICATION ===")
print("Excel Stock count:", len(excel_stock))

# Read realStockData.ts
with open(r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\frontend\src\data\realStockData.ts", "r", encoding="utf-8") as f:
    s_code = f.read()

stock_refs_found = re.findall(r"reference:\s*[\"'](PR-\d+)[\"']", s_code)
print(f"Found {len(stock_refs_found)} PR-xxx references in realStockData.ts")

missing_stock = set(excel_stock.keys()) - set(stock_refs_found)
if missing_stock:
    print("MISSING STOCK REFS IN FRONTEND:", sorted(list(missing_stock)))
else:
    print("ALL 30 Stock PR-xxx items exist in realStockData.ts!")

# 3. Preventive Sheet (8. Planning annuel)
sheet_prev = wb.worksheets[2]
excel_prev = []
for r in range(5, 42):
    eq_id = sheet_prev.cell(r, 1).value
    if eq_id:
        excel_prev.append({
            'equipId': eq_id,
            'equipName': sheet_prev.cell(r, 2).value,
            'tache': sheet_prev.cell(r, 3).value,
            'type': sheet_prev.cell(r, 4).value,
            'frequence': sheet_prev.cell(r, 5).value,
            'gammeRef': sheet_prev.cell(r, 6).value,
            'duree': float(sheet_prev.cell(r, 7).value or 0),
            'responsable': sheet_prev.cell(r, 8).value
        })

print("\n=== 3. PREVENTIVE TASKS VERIFICATION ===")
print("Excel Preventive tasks count:", len(excel_prev))
