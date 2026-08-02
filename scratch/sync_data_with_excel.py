import openpyxl
import json

excel_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\GMAO_FabLab_Universiapolis_Dossier_complet_8-9-12-13-14.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Equipments Sheet (14 machines)
sheet_eq = wb.worksheets[1]
excel_machines = []
for r in range(5, 19):
    eq_id = sheet_eq.cell(r, 1).value
    if eq_id:
        excel_machines.append({
            'id': eq_id,
            'reference': eq_id,
            'codeEq': f"EQ{len(excel_machines)+1}",
            'name': sheet_eq.cell(r, 2).value,
            'category': f"Atelier {sheet_eq.cell(r, 3).value}",
            'atelier': sheet_eq.cell(r, 3).value,
            'location': f"Espace {sheet_eq.cell(r, 3).value}",
            'manufacturer': sheet_eq.cell(r, 4).value or "FabLab Manufacturer",
            'model': sheet_eq.cell(r, 4).value or "Modèle Spécifique",
            'serialNumber': f"SN-{eq_id}-2023",
            'yearInService': 2021,
            'criticite': sheet_eq.cell(r, 5).value,
            'status': "Opérationnel",
            'quantity': 1,
            'description': f"Équipement principal {sheet_eq.cell(r, 2).value} - FabLab Universiapolis",
            'logiciel': "Logiciel de pilotage GMAO",
            'level': 1
        })

print(f"Loaded {len(excel_machines)} machines from Excel sheet 'Réf. Équipements'")

# 2. Stock Sheet (30 items: PR-001 to PR-030)
sheet_stk = wb.worksheets[4]
excel_stock = []
for r in range(5, 35):
    ref = sheet_stk.cell(r, 1).value
    if ref:
        excel_stock.append({
            'id': ref,
            'reference': ref,
            'name': sheet_stk.cell(r, 2).value,
            'category': sheet_stk.cell(r, 3).value,
            'zone': "Magasin GMAO",
            'subcat': sheet_stk.cell(r, 3).value,
            'equipement': sheet_stk.cell(r, 4).value,
            'supplier': sheet_stk.cell(r, 5).value,
            'refSupplier': sheet_stk.cell(r, 6).value or ref,
            'unitCostMAD': float(sheet_stk.cell(r, 7).value or 0),
            'priceMAD': float(sheet_stk.cell(r, 7).value or 0),
            'unit': sheet_stk.cell(r, 8).value or 'u',
            'quantity': int(float(sheet_stk.cell(r, 9).value or 0)),
            'min': int(float(sheet_stk.cell(r, 10).value or 0)),
            'max': int(float(sheet_stk.cell(r, 11).value or 0)),
            'location': sheet_stk.cell(r, 14).value or 'Casier Général',
            'description': f"Pièce de rechange {sheet_stk.cell(r, 2).value} pour {sheet_stk.cell(r, 4).value}",
            'status': "Disponible" if float(sheet_stk.cell(r, 9).value or 0) > 0 else "Rupture de Stock"
        })

print(f"Loaded {len(excel_stock)} stock items from Excel sheet '9. Pièces de rechange'")

# 3. Preventive Tasks Sheet (37 tasks)
sheet_prev = wb.worksheets[2]
excel_prev = []
for r in range(5, 42):
    eq_id = sheet_prev.cell(r, 1).value
    if eq_id:
        excel_prev.append({
            'id': f"PREV-{len(excel_prev)+1:03d}",
            'equipId': eq_id,
            'equipName': sheet_prev.cell(r, 2).value,
            'tache': sheet_prev.cell(r, 3).value,
            'type': sheet_prev.cell(r, 4).value,
            'frequence': sheet_prev.cell(r, 5).value,
            'gammeRef': sheet_prev.cell(r, 6).value,
            'duree': float(sheet_prev.cell(r, 7).value or 0),
            'responsable': sheet_prev.cell(r, 8).value,
            'jan': bool(sheet_prev.cell(r, 9).value),
            'feb': bool(sheet_prev.cell(r, 10).value),
            'mar': bool(sheet_prev.cell(r, 11).value),
            'apr': bool(sheet_prev.cell(r, 12).value),
            'may': bool(sheet_prev.cell(r, 13).value),
            'jun': bool(sheet_prev.cell(r, 14).value),
            'jul': bool(sheet_prev.cell(r, 15).value),
            'aug': bool(sheet_prev.cell(r, 16).value),
            'sep': bool(sheet_prev.cell(r, 17).value),
            'oct': bool(sheet_prev.cell(r, 18).value),
            'nov': bool(sheet_prev.cell(r, 19).value),
            'dec': bool(sheet_prev.cell(r, 20).value),
        })

print(f"Loaded {len(excel_prev)} preventive tasks from Excel sheet '8. Planning annuel'")
