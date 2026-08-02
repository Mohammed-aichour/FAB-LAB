import openpyxl

file_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\DOSSIER Tableau de criticité (A-B-C) - Plan de maintenance préventive complet -  Gammes de maintenance illustrées - FABLAB UNIVERSIAPOLIS.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

sheet = wb['02_Arborescence_Reelle']
print(f"=== Sheet: {sheet.title} ({sheet.max_row} rows) ===")

for r in range(1, sheet.max_row + 1):
    vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Row {r:02d}: {vals}")
