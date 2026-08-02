import openpyxl

file_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\DOSSIER Tableau de criticité (A-B-C) - Plan de maintenance préventive complet -  Gammes de maintenance illustrées - FABLAB UNIVERSIAPOLIS.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

for idx in range(0, min(5, len(wb.sheetnames))):
    sheet = wb.worksheets[idx]
    print(f"\n=== Sheet {idx}: {sheet.title} ===")
    for r in range(1, min(sheet.max_row + 1, 20)):
        vals = [sheet.cell(r, c).value for c in range(1, min(sheet.max_column + 1, 12))]
        if any(v is not None for v in vals):
            print(f"  Row {r}: {vals}")
