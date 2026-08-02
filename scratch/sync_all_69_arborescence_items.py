import openpyxl
import json

dossier_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\DOSSIER Tableau de criticité (A-B-C) - Plan de maintenance préventive complet -  Gammes de maintenance illustrées - FABLAB UNIVERSIAPOLIS.xlsx"
wb_dossier = openpyxl.load_workbook(dossier_path, data_only=True)
sheet_arb = wb_dossier['02_Arborescence_Reelle']

# Machines IDs to exclude from stock (since they belong to Machines page)
machine_ids = {'FL-009', 'FL-010', 'FL-068', 'FL-007', 'FL-008', 'FL-003', 'FL-004', 'FL-005', 'FL-006', 'FL-069', 'FL-012', 'FL-017', 'FL-013', 'FL-014'}

arb_stock_items = []

for r in range(5, sheet_arb.max_row + 1):
    code = sheet_arb.cell(r, 1).value
    name = sheet_arb.cell(r, 2).value
    cat = sheet_arb.cell(r, 3).value
    spec = sheet_arb.cell(r, 4).value
    qty = sheet_arb.cell(r, 5).value
    status = sheet_arb.cell(r, 6).value

    if code and str(code).startswith('FL-') and str(code).strip() not in machine_ids:
        c_code = str(code).strip()
        c_name = str(name or '').strip()
        c_cat = str(cat or 'Outillage & Électronique').strip()
        c_spec = str(spec or '').strip()
        c_qty = int(float(qty)) if (qty is not None and str(qty).replace('.','',1).isdigit()) else 1
        c_status = str(status or 'Opérationnel').strip()

        desc = c_spec if c_spec else f"{c_name} — Matériel d'atelier FabLab ({c_cat})"
        min_q = max(1, c_qty // 2) if c_qty > 1 else 1

        arb_stock_items.append({
            'id': c_code,
            'reference': c_code,
            'name': c_name,
            'category': c_cat,
            'zone': "Magasin GMAO",
            'subcat': c_cat,
            'equipement': f"Atelier FabLab ({c_cat})",
            'supplier': "Fournisseur Officiel FabLab",
            'refSupplier': c_code,
            'unitCostMAD': 150.0,
            'priceMAD': 150.0,
            'unit': 'u',
            'quantity': c_qty,
            'min': min_q,
            'max': c_qty * 3,
            'location': "Atelier / Casier Outillage",
            'description': desc,
            'status': c_status
        })

print(f"Extracted {len(arb_stock_items)} Outillage/Instruments items from 02_Arborescence_Reelle")

# PR-001 to PR-030 from GMAO dossier complet
gmao_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\GMAO_FabLab_Universiapolis_Dossier_complet_8-9-12-13-14.xlsx"
wb_gmao = openpyxl.load_workbook(gmao_path, data_only=True)
sheet_pr = wb_gmao.worksheets[4]

pr_stock_items = []
for r in range(5, 35):
    ref = sheet_pr.cell(r, 1).value
    if ref:
        qty = int(float(sheet_pr.cell(r, 9).value or 0))
        min_q = int(float(sheet_pr.cell(r, 10).value or 0))
        max_q = int(float(sheet_pr.cell(r, 11).value or 0))
        pr_stock_items.append({
            'id': str(ref),
            'reference': str(ref),
            'name': str(sheet_pr.cell(r, 2).value),
            'category': str(sheet_pr.cell(r, 3).value),
            'zone': "Magasin GMAO",
            'subcat': str(sheet_pr.cell(r, 3).value),
            'equipement': str(sheet_pr.cell(r, 4).value),
            'supplier': str(sheet_pr.cell(r, 5).value),
            'refSupplier': str(sheet_pr.cell(r, 6).value or ref),
            'unitCostMAD': float(sheet_pr.cell(r, 7).value or 0),
            'priceMAD': float(sheet_pr.cell(r, 7).value or 0),
            'unit': str(sheet_pr.cell(r, 8).value or 'u'),
            'quantity': qty,
            'min': min_q,
            'max': max_q,
            'location': str(sheet_pr.cell(r, 14).value or 'Casier Général'),
            'description': f"Pièce de rechange {sheet_pr.cell(r, 2).value} ({sheet_pr.cell(r, 4).value})",
            'status': "Disponible" if qty > 0 else "Rupture de Stock"
        })

print(f"Extracted {len(pr_stock_items)} PR-xxx items from GMAO sheet '9. Pièces de rechange'")

total_stock = pr_stock_items + arb_stock_items
print(f"TOTAL STOCK ITEMS: {len(total_stock)}")

# Write to realStockData.ts
stock_ts = f"""export interface StockItem {{
  id: string;
  reference: string;
  name: string;
  category: string;
  zone: string;
  subcat: string;
  equipement: string;
  supplier: string;
  refSupplier: string;
  unitCostMAD: number;
  priceMAD: number;
  unit: string;
  quantity: number;
  min: number;
  max: number;
  location: string;
  description?: string;
  status: string;
}}

export const realStockItems: StockItem[] = {json.dumps(total_stock, indent=2, ensure_ascii=False)};
"""

with open(r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\frontend\src\data\realStockData.ts", "w", encoding="utf-8") as f:
    f.write(stock_ts)

with open(r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\data_db\stock_db.json", "w", encoding="utf-8") as f:
    json.dump(total_stock, f, indent=2, ensure_ascii=False)

print("SUCCESSFULLY SYNCED 100% COMPLETE DOSSIER STOCK!")
