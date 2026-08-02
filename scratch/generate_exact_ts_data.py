import openpyxl
import json
import os

excel_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\GMAO_FabLab_Universiapolis_Dossier_complet_8-9-12-13-14.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Equipments Sheet (14 machines)
sheet_eq = wb.worksheets[1]
excel_machines = []
for r in range(5, 19):
    eq_id = sheet_eq.cell(r, 1).value
    if eq_id:
        excel_machines.append({
            'id': eq_id,
            'reference': eq_id,
            'codeEq': f"EQ{len(excel_machines)+1}",
            'name': sheet_eq.cell(r, 2).value,
            'category': f"Atelier {sheet_eq.cell(r, 3).value}",
            'atelier': sheet_eq.cell(r, 3).value,
            'location': f"Espace {sheet_eq.cell(r, 3).value}",
            'manufacturer': sheet_eq.cell(r, 4).value or "FabLab Manufacturer",
            'model': sheet_eq.cell(r, 4).value or "Modèle Spécifique",
            'serialNumber': f"SN-{eq_id}-2023",
            'yearInService': 2021,
            'criticite': sheet_eq.cell(r, 5).value,
            'status': "Opérationnel",
            'quantity': 1,
            'description': f"Équipement principal {sheet_eq.cell(r, 2).value} - FabLab Universiapolis",
            'logiciel': "Logiciel de pilotage GMAO",
            'level': 1
        })

machines_ts = f"""export interface Machine {{
  id: string | number;
  reference: string;
  codeEq: string;
  name: string;
  category: string;
  atelier: string;
  location: string;
  manufacturer: string;
  model: string;
  serialNumber: string;
  yearInService: number;
  criticite: 'A' | 'B' | 'C';
  status: 'Opérationnel' | 'Hors service' | 'À contrôler';
  quantity: number;
  description: string;
  logiciel: string;
  level: number;
}}

export const realMachinesData: Machine[] = {json.dumps(excel_machines, indent=2, ensure_ascii=False)};
"""

with open(r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\frontend\src\data\realMachineData.ts", "w", encoding="utf-8") as f:
    f.write(machines_ts)

# 2. Stock Sheet (30 items: PR-001 to PR-030)
sheet_stk = wb.worksheets[4]
excel_stock = []
for r in range(5, 35):
    ref = sheet_stk.cell(r, 1).value
    if ref:
        excel_stock.append({
            'id': ref,
            'reference': ref,
            'name': sheet_stk.cell(r, 2).value,
            'category': sheet_stk.cell(r, 3).value,
            'zone': "Magasin GMAO",
            'subcat': sheet_stk.cell(r, 3).value,
            'equipement': sheet_stk.cell(r, 4).value,
            'supplier': sheet_stk.cell(r, 5).value,
            'refSupplier': sheet_stk.cell(r, 6).value or ref,
            'unitCostMAD': float(sheet_stk.cell(r, 7).value or 0),
            'priceMAD': float(sheet_stk.cell(r, 7).value or 0),
            'unit': sheet_stk.cell(r, 8).value or 'u',
            'quantity': int(float(sheet_stk.cell(r, 9).value or 0)),
            'min': int(float(sheet_stk.cell(r, 10).value or 0)),
            'max': int(float(sheet_stk.cell(r, 11).value or 0)),
            'location': sheet_stk.cell(r, 14).value or 'Casier Général',
            'description': f"Pièce de rechange {sheet_stk.cell(r, 2).value} pour {sheet_stk.cell(r, 4).value}",
            'actionEntretien': "Contrôle visuel & usure périodique",
            'niveauRequis': "N1",
            'consigneSecurite': "Port des EPI obligatoires",
            'status': "Disponible" if float(sheet_stk.cell(r, 9).value or 0) > 0 else "Rupture de Stock"
        })

stock_ts = f"""export interface StockItem {{
  id: string;
  reference: string;
  name: string;
  category: string;
  zone: string;
  subcat: string;
  equipement: string;
  supplier: string;
  refSupplier: string;
  unitCostMAD: number;
  priceMAD: number;
  unit: string;
  quantity: number;
  min: number;
  max: number;
  location: string;
  description?: string;
  actionEntretien?: string;
  niveauRequis?: string;
  consigneSecurite?: string;
  status: string;
}}

export const realStockItems: StockItem[] = {json.dumps(excel_stock, indent=2, ensure_ascii=False)};
"""

with open(r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\frontend\src\data\realStockData.ts", "w", encoding="utf-8") as f:
    f.write(stock_ts)

# 3. Preventive Sheet (37 tasks)
sheet_prev = wb.worksheets[2]
excel_prev = []
for r in range(5, 42):
    eq_id = sheet_prev.cell(r, 1).value
    if eq_id:
        months_arr = []
        months_names = ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec']
        for idx, m_name in enumerate(months_names, start=9):
            if sheet_prev.cell(r, idx).value:
                months_arr.append(m_name)
        excel_prev.append({
            'id': f"PREV-{len(excel_prev)+1:03d}",
            'machineId': eq_id,
            'machineName': sheet_prev.cell(r, 2).value,
            'task': sheet_prev.cell(r, 3).value,
            'type': sheet_prev.cell(r, 4).value,
            'frequency': sheet_prev.cell(r, 5).value,
            'gammeRef': sheet_prev.cell(r, 6).value,
            'durationHours': float(sheet_prev.cell(r, 7).value or 0),
            'technician': sheet_prev.cell(r, 8).value,
            'months': months_arr,
            'status': 'Planifié',
            'lastDone': '2026-01-15'
        })

amdec_ts = f"""export interface PreventifTask {{
  id: string;
  machineId: string;
  machineName: string;
  task: string;
  type: string;
  frequency: string;
  gammeRef: string;
  durationHours: number;
  technician: string;
  months: string[];
  status: string;
  lastDone?: string;
}}

export interface AMDECItem {{
  id: string;
  machineName: string;
  component: string;
  failureMode: string;
  cause: string;
  effect: string;
  severity: number;
  occurrence: number;
  detection: number;
  rpn: number;
  criticite: 'A' | 'B' | 'C';
  action: string;
  frequency: string;
}}

export const preventifData: PreventifTask[] = {json.dumps(excel_prev, indent=2, ensure_ascii=False)};

export const amdecData: AMDECItem[] = [
  {{
    id: "AMD-001",
    machineName: "Imprimante 3D FDM (grand format)",
    component: "Extrudeur / Buse",
    failureMode: "Bouchage buse / extrusion sous-dimensionnée",
    cause: "Dépôt filament & surchauffe",
    effect: "Arrêt impression & pièce défectueuse",
    severity: 8,
    occurrence: 6,
    detection: 4,
    rpn: 192,
    criticite: "A",
    action: "Nettoyage hebdomadaire & buse de rechange",
    frequency: "Hebdomadaire"
  }},
  {{
    id: "AMD-002",
    machineName: "Découpe / gravure laser CO2",
    component: "Lentille / Miroir optique",
    failureMode: "Encrassement optique / perte de puissance",
    cause: "Fumées & résidus de découpe",
    effect: "Gravure incomplète & risque d'incendie",
    severity: 9,
    occurrence: 5,
    detection: 3,
    rpn: 135,
    criticite: "A",
    action: "Contrôle & nettoyage optique quotidien",
    frequency: "Quotidien"
  }},
  {{
    id: "AMD-003",
    machineName: "Fraiseuse CNC 3 axes (précision)",
    component: "Broche / Roulements",
    failureMode: "Surchauffe / jeu mécanique",
    cause: "Usure roulements & lubrification insuffisante",
    effect: "Perte de précision & usinage hors tolérance",
    severity: 9,
    occurrence: 4,
    detection: 4,
    rpn: 144,
    criticite: "A",
    action: "Vidange & contrôle vibration mensuel",
    frequency: "Mensuel"
  }}
];
"""

with open(r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\frontend\src\data\amdecData.ts", "w", encoding="utf-8") as f:
    f.write(amdec_ts)

print("FIXED interface properties & amdecData export!")
