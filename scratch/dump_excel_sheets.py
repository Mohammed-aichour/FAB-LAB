import openpyxl

excel_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\GMAO_FabLab_Universiapolis_Dossier_complet_8-9-12-13-14.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

def dump_sheet(sheet_name, max_r=25):
    sheet = wb[sheet_name]
    print(f"\n=================== {sheet_name} ===================")
    for r in range(1, min(sheet.max_row + 1, max_r)):
        row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        if any(v is not None for v in row_vals):
            print(f"Row {r}: {row_vals}")

dump_sheet('Rf. quipements', 20)
dump_sheet('9. Pices de rechange', 36)
dump_sheet('8. Planning annuel', 45)
