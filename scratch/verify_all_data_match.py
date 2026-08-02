import openpyxl
import json
import re

excel_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\GMAO_FabLab_Universiapolis_Dossier_complet_8-9-12-13-14.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Equipments
sheet_eq = wb.worksheets[1]
eq_excel = []
for r in range(5, 19):
    eq_excel.append({
        'id': sheet_eq.cell(r, 1).value,
        'name': sheet_eq.cell(r, 2).value,
        'atelier': sheet_eq.cell(r, 3).value,
        'manufacturer': sheet_eq.cell(r, 4).value,
        'criticite': sheet_eq.cell(r, 5).value,
    })

print(f"Excel Equipments count: {len(eq_excel)}")

# 2. Stock Items
sheet_stk = wb.worksheets[4]
stk_excel = []
for r in range(5, 35):
    stk_excel.append({
        'ref': sheet_stk.cell(r, 1).value,
        'name': sheet_stk.cell(r, 2).value,
        'category': sheet_stk.cell(r, 3).value,
        'equipement': sheet_stk.cell(r, 4).value,
        'supplier': sheet_stk.cell(r, 5).value,
        'priceMAD': sheet_stk.cell(r, 7).value,
        'stock': sheet_stk.cell(r, 9).value,
        'min': sheet_stk.cell(r, 10).value,
    })

print(f"Excel Stock items count: {len(stk_excel)}")

# 3. Preventive Tasks
sheet_prev = wb.worksheets[2]
prev_excel = []
for r in range(5, 42):
    prev_excel.append({
        'fl_id': sheet_prev.cell(r, 1).value,
        'equipement': sheet_prev.cell(r, 2).value,
        'tache': sheet_prev.cell(r, 3).value,
        'type': sheet_prev.cell(r, 4).value,
        'frequence': sheet_prev.cell(r, 5).value,
        'gamme': sheet_prev.cell(r, 6).value,
        'duree': sheet_prev.cell(r, 7).value,
        'responsable': sheet_prev.cell(r, 8).value,
    })

print(f"Excel Preventive tasks count: {len(prev_excel)}")
