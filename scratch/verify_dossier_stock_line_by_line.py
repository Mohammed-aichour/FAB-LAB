import openpyxl
import json

file_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\DOSSIER Tableau de criticité (A-B-C) - Plan de maintenance préventive complet -  Gammes de maintenance illustrées - FABLAB UNIVERSIAPOLIS.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)
sheet = wb['02_Arborescence_Reelle']

print("=== DOSSIER 02_Arborescence_Reelle Verification ===")

excel_rows = []
for r in range(4, sheet.max_row + 1):
    code = sheet.cell(r, 1).value
    name = sheet.cell(r, 2).value
    type_or_cat = sheet.cell(r, 3).value
    model_spec = sheet.cell(r, 4).value
    qty = sheet.cell(r, 5).value
    status = sheet.cell(r, 6).value
    
    if code and str(code).startswith('FL-'):
        excel_rows.append({
            'code': str(code).strip(),
            'name': str(name or '').strip(),
            'category': str(type_or_cat or '').strip(),
            'spec': str(model_spec or '').strip(),
            'quantity': int(float(qty)) if (qty is not None and str(qty).replace('.','',1).isdigit()) else 1,
            'status': str(status or '').strip()
        })

print(f"Total FL-xxx rows found in 02_Arborescence_Reelle: {len(excel_rows)}")

# Read realStockData.ts
with open(r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\frontend\src\data\realStockData.ts", "r", encoding="utf-8") as f:
    stock_ts_text = f.read()

# Read realMachineData.ts
with open(r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\frontend\src\data\realMachineData.ts", "r", encoding="utf-8") as f:
    machine_ts_text = f.read()

found_in_stock = 0
found_in_machines = 0
missing_anywhere = []

for row in excel_rows:
    c = row['code']
    in_stock = c in stock_ts_text
    in_mach = c in machine_ts_text
    
    if in_stock:
        found_in_stock += 1
    if in_mach:
        found_in_machines += 1
        
    if not in_stock and not in_mach:
        missing_anywhere.append(row)

print(f"Items present in Machines: {found_in_machines}")
print(f"Items present in Stock: {found_in_stock}")
print(f"Items missing anywhere: {len(missing_anywhere)}")

if missing_anywhere:
    print("\n--- MISSING ITEMS DETAILS ---")
    for item in missing_anywhere:
        print(f"Code: {item['code']} | Name: {item['name']} | Qty: {item['quantity']} | Spec: {item['spec']}")
