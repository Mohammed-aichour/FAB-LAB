import openpyxl

excel_path = r"d:\Apps\Desktop\Downloads\Stage_3A\Anti\GMAO_FabLab_Universiapolis_Dossier_complet_8-9-12-13-14.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

def dump_sheet_by_idx(idx, max_r=45):
    sheet = wb.worksheets[idx]
    print(f"\n=================== Sheet {idx}: {sheet.title} ===================")
    for r in range(1, min(sheet.max_row + 1, max_r)):
        row_vals = [sheet.cell(r, c).value for c in range(1, min(sheet.max_column + 1, 15))]
        if any(v is not None for v in row_vals):
            print(f"Row {r}: {row_vals}")

dump_sheet_by_idx(1, 25) # Equipements
dump_sheet_by_idx(2, 45) # Planning annuel
dump_sheet_by_idx(4, 36) # Pieces de rechange
